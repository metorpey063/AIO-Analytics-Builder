"""
Tableau Studio skills builder for AIO Analytics Builder.

Generates Tableau Studio Agent skills pre-loaded with a demo's story, so the SE or the audience can
build and explain the demo in natural language inside Tableau Studio:

  aio-story         guided builder: walks the four-question narrative (What's wrong → Where → Why → Who)
  aio-autobuilder   one prompt ("build the demo") creates the Q1–Q4 worksheets and the story dashboard
  aio-audience      interpreter for presenters and viewers ("why did this move?")
  + any custom skills Claude writes for the demo (validated and installed alongside)

Templates live in sidekicks/studio/skills/. Output lands in demos/{slug}/studio_skills/.
The story is computed from the data with close_loop.analyze_story, so it matches what close-loop tells.

CLI:
  python studio_skills.py story    --slug <slug> [--csv PATH | --hyper PATH | --datasource NAME] [--new] [--set key=value ...]
  python studio_skills.py generate --slug <slug>
  python studio_skills.py custom   --slug <slug> --name <kebab-name> --spec SKILL.md [--refs DIR]
  python studio_skills.py install  --slug <slug>
  python studio_skills.py status   --slug <slug>
"""

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime

import pandas as pd

from close_loop import (ROOT, _change, _find_csv, _is_avg, _literal_assignments, analyze_story, cp_path,
                        load_cp, save_cp)

TEMPLATES = os.path.join(ROOT, "sidekicks", "studio", "skills")
TEMPLATE_SKILLS = ["aio-story", "aio-audience", "aio-autobuilder"]
STUDIO_SKILLS_DIR = os.path.expanduser("~/Documents/My Tableau Repository (Beta)/Tableau Agent/skills")
SECTION_TITLE = "Tableau Studio Skills"
DEFAULT_LAG = 0.7
NAME_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")


# ── Paths ──────────────────────────────────────────────────────────────────────

def demo_path(slug, create=False):
    d = os.path.join(ROOT, "demos", slug)
    if create:
        os.makedirs(d, exist_ok=True)
    elif not os.path.isdir(d):
        raise SystemExit(f"Demo folder not found: {d} — pass --csv, --hyper, or --datasource to create it.")
    return d


def out_dir(slug):
    return os.path.join(demo_path(slug), "studio_skills")


def _rel(path):
    return os.path.relpath(path, ROOT) if os.path.isabs(path) and path.startswith(ROOT) else path


# ── Sources: external CSV / .hyper / published datasource ─────────────────────

def _hyper_to_csv(hyper_path, dest_csv):
    """Read the largest table in a .hyper into a CSV."""
    from tableauhyperapi import Connection, HyperProcess, Telemetry
    log_dir = tempfile.mkdtemp(prefix="aio_hyper_")
    with HyperProcess(telemetry=Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU,
                      parameters={"log_dir": log_dir}) as hp:
        with Connection(hp.endpoint, hyper_path) as conn:
            tables = [t for s in conn.catalog.get_schema_names() for t in conn.catalog.get_table_names(s)]
            if not tables:
                raise SystemExit(f"No tables found in {hyper_path}")
            table = max(tables, key=lambda t: conn.execute_scalar_query(f"SELECT COUNT(*) FROM {t}"))
            cols = [c.name.unescaped for c in conn.catalog.get_table_definition(table).columns]
            rows = conn.execute_list_query(f"SELECT * FROM {table}")

    def py(v):  # Hyper Date / Timestamp → Python
        for attr in ("to_datetime", "to_date"):
            if hasattr(v, attr):
                return getattr(v, attr)()
        return v
    df = pd.DataFrame([[py(v) for v in r] for r in rows], columns=cols)
    if "Day_Offset" in df and not any(c.lower() in ("date", "week", "month") for c in df.columns):
        # Self-healing .tdsx: Date is a calc (DATEADD('day', [Day_Offset], TODAY())), not stored
        df.insert(1, "Date", (pd.Timestamp.today().normalize() + pd.to_timedelta(df["Day_Offset"], unit="D")).dt.date)
    df.to_csv(dest_csv, index=False)
    shutil.rmtree(log_dir, ignore_errors=True)
    return dest_csv


def _hyper_from_tdsx(tdsx_path, dest_dir):
    with zipfile.ZipFile(tdsx_path) as z:
        hypers = [n for n in z.namelist() if n.endswith(".hyper")]
        if not hypers:
            raise SystemExit(f"{os.path.basename(tdsx_path)} has no extract (.hyper) — can't read its data.")
        z.extract(hypers[0], dest_dir)
        return os.path.join(dest_dir, hypers[0])


def _download_datasource(name, dest_dir, profile=None):
    import tableauserverclient as TSC
    from connections import get_tableau_token, load_config
    server, _, _ = get_tableau_token(load_config(profile) if profile else None)
    try:
        opts = TSC.RequestOptions()
        opts.filter.add(TSC.Filter(TSC.RequestOptions.Field.Name, TSC.RequestOptions.Operator.Equals, name))
        matches, _ = server.datasources.get(opts)
        if not matches:
            raise SystemExit(f"No published datasource named '{name}' on this site.")
        return server.datasources.download(matches[0].id, filepath=dest_dir, include_extract=True)
    finally:
        server.auth.sign_out()


def ingest_source(slug, csv=None, hyper=None, datasource=None, profile=None):
    """Bring outside demo content into demos/{slug}/ as a CSV and record it on the checkpoint story."""
    d = demo_path(slug, create=True)
    cp = load_cp(slug)
    story = cp.get("story") or {}
    if csv:
        dest = os.path.join(d, os.path.basename(csv))
        if os.path.abspath(csv) != os.path.abspath(dest):
            shutil.copy(csv, dest)
    else:
        if datasource:
            src = _download_datasource(datasource, d, profile)
            story["published_datasource"] = datasource
            hyper = src
        tmp = tempfile.mkdtemp(prefix="aio_tdsx_")
        if hyper.endswith(".tdsx"):
            hyper = _hyper_from_tdsx(hyper, tmp)
        dest = _hyper_to_csv(hyper, os.path.join(d, f"{slug}_data.csv"))
        shutil.rmtree(tmp, ignore_errors=True)
    story["datasource"] = _rel(dest)
    cp["story"] = story
    cp["csv_path"] = dest
    save_cp(slug, cp)
    print(f"✓ Source ready: {_rel(dest)}")
    return dest


# ── Story ──────────────────────────────────────────────────────────────────────

def _context(slug, cp, story):
    scripts = glob.glob(os.path.join(demo_path(slug), "*_demo.py"))
    lits = _literal_assignments(scripts[0] if scripts else None,
                                {"COMPANY", "USE_CASE_LABEL", "USE_CASE", "METRIC_CONFIG", "SIGNAL_ONSET",
                                 "SIGNAL_SHAPE"})
    mc = lits.get("METRIC_CONFIG") or []
    if not mc and story.get("primary_metric"):  # outside demo: metrics come from the story
        aggs = story.get("aggregations") or {}
        labels = [story["primary_metric"]] + list(story.get("supporting_metrics") or [])
        mc = [{"label": m, "field": m, "agg": aggs.get(m, "AVG"),
               "sentiment": story.get("sentiment", "UP_IS_GOOD")} for m in labels]
    return lits, {
        "company": story.get("company") or lits.get("COMPANY") or slug.replace("_", " ").title(),
        "use_case": lits.get("USE_CASE_LABEL") or lits.get("USE_CASE") or "",
        "metric_config": mc,
        "story": story,
        "csv_path": _find_csv(slug, cp),
    }


def _mult(seg, overall):
    if not seg or not overall or not overall.get("pct"):
        return None
    return round(seg["pct"] / overall["pct"], 2)


def _pct(x):
    return f"{x['pct'] * 100:+.0f}%" if x else "n/a"


def _parse_value(v):
    try:
        return json.loads(v)
    except (ValueError, TypeError):
        return v


def build_story(slug, new=False, sets=None):
    """Fill the story from the data. Saved keys win over derived ones unless new=True; --set wins over both."""
    cp = load_cp(slug)
    saved = {} if new else dict(cp.get("story") or {})
    for kv in sets or []:
        k, _, v = kv.partition("=")
        saved[k.strip()] = _parse_value(v.strip())
    if new:  # keep where the data lives, re-derive everything else
        for k in ("datasource", "published_datasource"):
            if (cp.get("story") or {}).get(k):
                saved.setdefault(k, cp["story"][k])

    lits, ctx = _context(slug, cp, saved)
    df = pd.read_csv(ctx["csv_path"])
    a = analyze_story(df, ctx)
    overall, culprit, compound, counter = a["overall"], a["culprit"], a["compound"], a["counter_trend"]

    # Moderate = the other values of the culprit dimension (same windows as analyze_story)
    date_field = saved.get("date_field") or next((c for c in df.columns if c.lower() in ("date", "week", "month")), "Date")
    df["_month"] = pd.to_datetime(df[date_field]).dt.to_period("M")
    latest = df["_month"].max()
    moderate = []
    for v in sorted({str(v) for v in df[culprit["dimension"]].dropna().unique()}):
        if v in (culprit["value"], (counter or {}).get("value")):
            continue
        seg = df[df[culprit["dimension"]].astype(str) == v]
        ch = _change(seg[seg["_month"] <= latest - 9], seg[seg["_month"] > latest - 3],
                     a["primary_field"], a["is_rate"])
        if ch:
            moderate.append({"name": v, "dimension": culprit["dimension"], "multiplier": _mult(ch, overall)})

    metric_labels = [m.get("label") for m in ctx["metric_config"]]
    aggs = {m["label"]: ("AVG" if _is_avg(m) else "SUM") for m in ctx["metric_config"] if m.get("label")}
    supporting = [s["label"] for s in a["supporting"]] or metric_labels[1:4]
    onset = lits.get("SIGNAL_ONSET")
    payload = _close_loop_payload(slug)
    direction = "down" if (overall or {}).get("pct", 0) < 0 else "up"

    derived = {
        "company": ctx["company"],
        "use_case": ctx["use_case"],
        "headline": (payload or {}).get("headline") or
                    f"{a['primary_metric']} {direction} {abs((overall or {}).get('pct', 0)) * 100:.0f}% — "
                    f"concentrated in {culprit['value']}" + (f" / {compound['value']}" if compound else ""),
        "primary_metric": a["primary_metric"],
        "supporting_metrics": supporting,  # labels, leading indicator first (shared with close-loop)
        "lag_ratios": {m: DEFAULT_LAG for m in supporting},
        "dimensions": [d for d in (saved.get("dimensions") or []) if d in df.columns] or
                      sorted({s["dimension"] for s in [culprit, compound, counter] if s} |
                             ({a["root_cause_dimension"]} if a["root_cause_dimension"] else set())),
        "culprit_segments": [{"name": culprit["value"], "dimension": culprit["dimension"],
                              "multiplier": _mult(culprit, overall)}],
        "compound_segment": ({"name": compound["value"], "dimension": compound["dimension"],
                              "multiplier": _mult(compound, overall)} if compound else None),
        "moderate_segments": moderate,
        "counter_trend_segment": ({"name": counter["value"], "dimension": counter["dimension"],
                                   "multiplier": _mult(counter, overall)} if counter else None),
        "root_cause_dimension": a["root_cause_dimension"],
        "entity_field": a["entity_field"],
        "date_field": date_field,
        "aggregations": aggs or {a["primary_metric"]: "AVG" if a["is_rate"] else "SUM"},
        "signal_onset_months_ago": abs(onset) if isinstance(onset, (int, float)) else 3,
        "signal_shape": lits.get("SIGNAL_SHAPE") or "accelerating",
        "datasource": _rel(ctx["csv_path"]),
        "published_datasource": cp.get("pulse_ds_name"),
    }
    story = {**derived, **{k: v for k, v in saved.items() if v not in (None, "", [], {})}}
    story["evidence"] = {  # always recomputed from the data
        "overall": _pct(overall), "culprit": _pct(culprit), "compound": _pct(compound),
        "counter_trend": _pct(counter),
        "supporting": {s["label"]: _pct(s) for s in a["supporting"]},
        "at_risk": [e["entity"] for e in a["at_risk"]],
        "windows": a["windows"],
    }
    story = {k: v for k, v in story.items() if v is not None}
    cp = load_cp(slug)
    cp["story"] = story
    save_cp(slug, cp)
    return story


def _close_loop_payload(slug):
    p = os.path.join(demo_path(slug), f"{slug}_close_loop_payload.json")
    if os.path.exists(p):
        with open(p) as f:
            return json.load(f)
    return None


def story_summary(story):
    c = (story.get("culprit_segments") or [{}])[0]
    comp, ct, ev = story.get("compound_segment"), story.get("counter_trend_segment") or {}, story.get("evidence", {})
    lines = [
        f"Company:        {story.get('company')}",
        f"Headline:       {story.get('headline')}",
        f"Primary metric: {story.get('primary_metric')} ({ev.get('overall', 'n/a')} overall)",
        f"Culprit:        {c.get('dimension')} = {c.get('name')} ({ev.get('culprit', 'n/a')})",
    ]
    if comp:
        lines.append(f"Deepest drill:  + {comp['dimension']} = {comp['name']} ({ev.get('compound', 'n/a')})")
    lines += [
        f"Counter-trend:  {ct.get('dimension')} = {ct.get('name')} ({ev.get('counter_trend', 'n/a')})",
        f"Root cause dim: {story.get('root_cause_dimension')}",
        f"Supporting:     {', '.join(story.get('supporting_metrics') or []) or '—'}",
        f"Entity (Q4):    {story.get('entity_field')}   Date: {story.get('date_field')}",
        f"Datasource:     {story.get('published_datasource') or story.get('datasource')}",
    ]
    return "\n".join(lines)


# ── demo-story.md ──────────────────────────────────────────────────────────────

def render_demo_story(story, payload=None):
    ev = story.get("evidence", {})
    culprits = story.get("culprit_segments") or []
    comp, ct = story.get("compound_segment"), story.get("counter_trend_segment")
    lag = story.get("lag_ratios") or {}
    aggs = story.get("aggregations") or {}
    L = [f"# Demo Story — {story.get('company')}", "",
         f"**Headline:** {story.get('headline')}", "",
         "## Primary Metric", f"{story.get('primary_metric')} (aggregation: {aggs.get(story.get('primary_metric'), 'AVG')}). "
         f"Overall change, recent 3 months vs. baseline: {ev.get('overall', 'n/a')}.", "",
         "## The Problem",
         f"{story.get('primary_metric')} has moved the wrong way since about {story.get('signal_onset_months_ago', 3)} "
         f"months ago ({story.get('signal_shape', 'accelerating')} shape: slow at first, steeper toward today). "
         "The overall average hides how concentrated the problem is.", "",
         "## Culprit Segment(s)"]
    for c in culprits:
        L.append(f"- **{c['name']}** ({c['dimension']}): {ev.get('culprit', 'n/a')}, "
                 f"about {c.get('multiplier') or '?'}× the overall move")
    if comp:
        L.append(f"- Deepest drill: **{culprits[0]['name'] if culprits else ''} + {comp['name']}** "
                 f"({comp['dimension']}): {ev.get('compound', 'n/a')}")
    L += ["", "## Moderate Segments"]
    L += [f"- {m['name']} ({m['dimension']}): about {m.get('multiplier')}× the overall move"
          for m in story.get("moderate_segments") or []] or ["- None identified"]
    L += ["", "## Counter-Trend Segment"]
    L.append(f"- **{ct['name']}** ({ct['dimension']}): {ev.get('counter_trend', 'n/a')}. This is the benchmark that "
             "proves the problem isn't systemic." if ct else "- None identified")
    L += ["", "## Root Cause Dimension",
          f"{story.get('root_cause_dimension')}: use it to explain *why* the culprit is worst.", "",
          "## Supporting Metrics (Leading Indicators)"]
    sup_ev = ev.get("supporting", {})
    L += [f"- {m}: {sup_ev.get(m, 'n/a')} in the culprit segment; lags the primary at about {lag.get(m, DEFAULT_LAG)}×"
          for m in story.get("supporting_metrics") or []] or ["- None"]
    L += ["", "## Dimensions for Drilling"] + [f"- {d}" for d in story.get("dimensions") or []]
    if ev.get("at_risk"):
        L += ["", "## At-Risk Entities (Q4)"] + [f"- {e}" for e in ev["at_risk"]]
    L += ["", "## Build Fields (used by aio-autobuilder)",
          f"- Headline: {story.get('headline')}",
          f"- Date field: {story.get('date_field')}",
          f"- Segment dimension (Q2): {culprits[0]['dimension'] if culprits else ''}",
          f"- Culprit value: {culprits[0]['name'] if culprits else ''}",
          f"- Counter-trend value: {ct['name'] if ct else ''}",
          f"- Root-cause dimension (Q3): {story.get('root_cause_dimension')}",
          f"- Entity field (Q4): {story.get('entity_field')}",
          "- Aggregations: " + ", ".join(f"{k} → {v}" for k, v in aggs.items()),
          f"- Datasource to connect: {story.get('published_datasource') or story.get('datasource')}"
          + (f" (or the CSV at {story['datasource']})" if story.get("published_datasource") else "")]
    if payload:
        L += ["", "## Close the Loop",
              f"A follow-up is installed in the Salesforce org: **{payload.get('action')}** — {payload.get('action_reason', '')}".rstrip(" —"),
              f"Trigger: {payload.get('trigger_label') or payload.get('trigger')}", ""]
        L += [f"- {s['title']}: \"{s['say']}\" — {s['trigger']} → {s['result']}" for s in payload.get("demo_steps", [])]
    return "\n".join(L) + "\n"


# ── Generate ───────────────────────────────────────────────────────────────────

def _skill_list(cp):
    return TEMPLATE_SKILLS + [s["name"] for s in cp.get("studio_custom_skills", [])]


def write_install_sh(slug, cp):
    company = (cp.get("story") or {}).get("company", slug)
    skills = " ".join(_skill_list(cp))
    script = f"""#!/bin/bash
# Install this demo's Tableau Studio Agent skills ({company} story). Generated by studio_skills.py.
set -e

SKILLS_DIR="$HOME/Documents/My Tableau Repository (Beta)/Tableau Agent/skills"
SCRIPT_DIR="$( cd "$( dirname "${{BASH_SOURCE[0]}}" )" && pwd )"

if [ ! -d "$SKILLS_DIR" ]; then
    echo "Error: Tableau Agent skills directory not found at:"
    echo "  $SKILLS_DIR"
    echo ""
    echo "Make sure you have:"
    echo "  1. Installed Tableau Studio"
    echo "  2. Opened Tableau Studio at least once"
    echo "  3. Opened Tableau Agent (Extensions menu)"
    exit 1
fi

for skill in {skills}; do
    if [ ! -d "$SCRIPT_DIR/$skill" ]; then
        echo "✗ $skill not found at $SCRIPT_DIR/$skill"
        exit 1
    fi
    rm -rf "$SKILLS_DIR/$skill"
    cp -r "$SCRIPT_DIR/$skill" "$SKILLS_DIR/$skill"
    echo "✓ Installed $skill ({company} story)"
done

echo ""
echo "In Tableau Agent, type: /reload-skills"
"""
    p = os.path.join(out_dir(slug), "install.sh")
    with open(p, "w") as f:
        f.write(script)
    os.chmod(p, 0o755)
    return p


def skill_prompts(story, cp):
    """(skill, prompts, what it does) — used by README.md and walkthrough Section 8."""
    c = ((story.get("culprit_segments") or [{}])[0]).get("name", "the culprit segment")
    ct = (story.get("counter_trend_segment") or {}).get("name", "the counter-trend segment")
    m, co, ent = story.get("primary_metric"), story.get("company"), story.get("entity_field", "account")
    out = [
        ("aio-autobuilder", ["build the demo"],
         f"Builds the four story worksheets (Q1 {m} trend, Q2 by {((story.get('culprit_segments') or [{}])[0]).get('dimension')} "
         f"with {c} in red and {ct} in orange, Q3 the leading indicators, Q4 the at-risk {ent} list) and a 2×2 "
         f"dashboard titled \"{co} — {story.get('headline')}\", checking each sheet as it goes."),
        ("aio-story", [f"Walk me through the AiO story for {co}", f"Build Q1: what's wrong with {m}",
                       f"Build Q2: where is it worst", "Build Q3: why is it happening",
                       f"Build Q4: which {ent} values need action", "Create the dashboard with all four questions"],
         f"Guides the presenter one question at a time, starting from this demo's story ({m}, culprit {c}), "
         "and explains the reason for each chart choice."),
        ("aio-audience", [f"Explain this {m} trend", f"Why is {c} down so much?",
                          f"Which {c} accounts are most at risk?", "What should we do about this?"],
         f"Tells the story in business language: {c} is driving the decline ({story.get('evidence', {}).get('culprit', '')}), "
         f"{ct} proves it isn't systemic, and {story.get('root_cause_dimension')} explains why"
         + (". It points to the Close the Loop follow-up when asked what to do." if _close_loop_payload_for(cp) else ".")),
    ]
    for s in cp.get("studio_custom_skills", []):
        out.append((s["name"], s.get("prompts") or [], s.get("description", "")))
    return out


def _close_loop_payload_for(cp):
    return bool(cp.get("close_loop"))


def write_readme(slug, story, cp):
    L = [f"# Tableau Studio skills — {story.get('company')}", "", f"_{story.get('headline')}_", "",
         "## Quick start",
         "1. `bash install.sh`",
         "2. In Tableau Agent, type `/reload-skills`",
         f"3. Connect the datasource: **{story.get('published_datasource') or story.get('datasource')}**",
         "4. Try a prompt below", ""]
    for skill, prompts, what in skill_prompts(story, cp):
        L += [f"## {skill}", what, ""] + [f"- \"{p}\"" for p in prompts] + [""]
    L += ["## More",
          "- This demo's story: `references/demo-story.md` in each skill",
          "- Design: `sidekicks/studio/docs/studio-skills-design.md`",
          "- Generic skills (no demo): `sidekicks/studio/skills/`, installed with `bash sidekicks/studio/install_studio_skills.sh`"]
    p = os.path.join(out_dir(slug), "README.md")
    with open(p, "w") as f:
        f.write("\n".join(L) + "\n")
    return p


def generate(slug):
    cp = load_cp(slug)
    story = cp.get("story")
    if not story or not story.get("culprit_segments"):
        story = build_story(slug)
        cp = load_cp(slug)
    md = render_demo_story(story, _close_loop_payload(slug))
    base = out_dir(slug)
    os.makedirs(base, exist_ok=True)
    for skill in TEMPLATE_SKILLS:
        dst = os.path.join(base, skill)
        shutil.rmtree(dst, ignore_errors=True)
        shutil.copytree(os.path.join(TEMPLATES, skill), dst)
    for skill in _skill_list(cp):  # custom skills keep their SKILL.md; only the story is refreshed
        refs = os.path.join(base, skill, "references")
        if os.path.isdir(os.path.join(base, skill)):
            os.makedirs(refs, exist_ok=True)
            with open(os.path.join(refs, "demo-story.md"), "w") as f:
                f.write(md)
    write_install_sh(slug, cp)
    write_readme(slug, story, cp)
    cp["studio_skills_generated"] = True
    cp.setdefault("studio_skills_installed", False)
    cp["studio_skills_dir"] = _rel(base)
    save_cp(slug, cp)
    doc = write_walkthrough(slug, cp, story)
    print(f"✓ Studio skills written to {_rel(base)}/ ({', '.join(_skill_list(cp))})")
    if doc:
        print(f"✓ Walkthrough Section 8 ({SECTION_TITLE}) written to {_rel(doc)}")
    return base


# ── Custom skills ──────────────────────────────────────────────────────────────

def _frontmatter(text):
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return None
    fm = {}
    for line in m.group(1).splitlines():
        k, sep, v = line.partition(":")
        if sep:
            fm[k.strip()] = v.strip().strip("\"'")
    return fm


def _reserved_names():
    names = set(TEMPLATE_SKILLS)
    for sub in ("overridable", "protected"):
        d = os.path.join(STUDIO_SKILLS_DIR, ".vendored", sub)
        if os.path.isdir(d):
            names |= {n for n in os.listdir(d) if not n.startswith(".")}
    return names


def add_custom_skill(slug, name, spec_path, refs_dir=None):
    """Validate a Claude-authored SKILL.md, add it to studio_skills/, and include it in install.sh."""
    if not NAME_RE.match(name) or len(name) > 64:
        raise SystemExit(f"Skill name '{name}' must be kebab-case (a-z, 0-9, hyphens), 64 chars max.")
    if name in _reserved_names():
        raise SystemExit(f"'{name}' clashes with an AiO template or a standard Tableau Studio skill — pick another name.")
    with open(spec_path) as f:
        text = f.read()
    fm = _frontmatter(text)
    if not fm:
        raise SystemExit("SKILL.md must start with YAML frontmatter (--- name / description ---).")
    if fm.get("name") != name:
        raise SystemExit(f"Frontmatter name '{fm.get('name')}' must match --name '{name}'.")
    if len(fm.get("description", "")) < 20:
        raise SystemExit("Frontmatter description is required: say what the skill does and list its trigger prompts.")

    cp = load_cp(slug)
    if not cp.get("studio_skills_generated"):
        generate(slug)
        cp = load_cp(slug)
    story = cp.get("story") or {}
    warnings = []
    if "demo-story.md" not in text:
        warnings.append("SKILL.md never mentions references/demo-story.md — it won't use this demo's story.")
    cols = set(pd.read_csv(_find_csv(slug, cp), nrows=0).columns)
    values = {str(s.get("name")) for s in (story.get("culprit_segments") or []) + (story.get("moderate_segments") or [])
              + [story.get("compound_segment") or {}, story.get("counter_trend_segment") or {}]}
    ticked = set(re.findall(r"`([A-Z][^`\n]{1,40})`", text))
    unknown = sorted(t for t in ticked if t not in cols and t not in values and "/" not in t and "." not in t)
    if unknown:
        warnings.append(f"Backticked names not found in the datasource or story: {', '.join(unknown)}")

    dst = os.path.join(out_dir(slug), name)
    shutil.rmtree(dst, ignore_errors=True)
    os.makedirs(os.path.join(dst, "references"))
    shutil.copy(spec_path, os.path.join(dst, "SKILL.md"))
    if refs_dir:
        for p in glob.glob(os.path.join(refs_dir, "*")):
            shutil.copy(p, os.path.join(dst, "references"))
    with open(os.path.join(dst, "references", "demo-story.md"), "w") as f:
        f.write(render_demo_story(story, _close_loop_payload(slug)))

    prompts = re.findall(r"\"([^\"]{4,80})\"", fm.get("description", ""))
    customs = [s for s in cp.get("studio_custom_skills", []) if s["name"] != name]
    customs.append({"name": name, "description": fm["description"], "prompts": prompts,
                    "created_at": datetime.now().isoformat(timespec="seconds")})
    cp["studio_custom_skills"] = customs
    cp["studio_skills_installed"] = False
    save_cp(slug, cp)
    write_install_sh(slug, cp)
    write_readme(slug, story, cp)
    write_walkthrough(slug, cp, story)
    print(f"✓ Custom skill {name} added to {_rel(dst)}/ and install.sh")
    for w in warnings:
        print(f"⚠ {w}")
    return dst


# ── Install / status ───────────────────────────────────────────────────────────

def install(slug):
    sh = os.path.join(out_dir(slug), "install.sh")
    if not os.path.exists(sh):
        generate(slug)
    r = subprocess.run(["bash", sh])
    if r.returncode != 0:
        raise SystemExit(r.returncode)
    cp = load_cp(slug)
    cp["studio_skills_installed"] = True
    save_cp(slug, cp)


def status(slug):
    cp = load_cp(slug)
    base = out_dir(slug)
    rep = {"generated": bool(cp.get("studio_skills_generated")), "dir": _rel(base),
           "studio_found": os.path.isdir(STUDIO_SKILLS_DIR), "skills": []}
    for s in _skill_list(cp):
        rep["skills"].append({"name": s, "generated": os.path.isdir(os.path.join(base, s)),
                              "installed": os.path.isdir(os.path.join(STUDIO_SKILLS_DIR, s)),
                              "custom": s not in TEMPLATE_SKILLS})
    print(json.dumps(rep, indent=2))
    return rep


# ── Walkthrough Section 8 ──────────────────────────────────────────────────────

def write_walkthrough(slug, cp, story):
    """Replace (or append) the 'Tableau Studio Skills' section in the demo's walkthrough .docx."""
    path = cp.get("walkthrough") or next(iter(sorted(glob.glob(os.path.join(demo_path(slug), "*walkthrough*.docx")))), None)
    if not path or not os.path.exists(path) or os.path.basename(path).startswith("~$"):
        return None
    from docx import Document
    doc = Document(path)
    drop = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            drop = p.text.strip() == SECTION_TITLE
        if drop:
            p._element.getparent().remove(p._element)
    doc.add_heading(SECTION_TITLE, level=1)
    doc.add_paragraph(
        "These Tableau Studio Agent skills already know this demo's story, so you or the audience can build and "
        "explain it in plain language inside Tableau Studio. To install, run "
        f"bash {_rel(os.path.join(out_dir(slug), 'install.sh'))}, type /reload-skills in Tableau Agent, then connect "
        f"the datasource ({story.get('published_datasource') or story.get('datasource')}).")
    for skill, prompts, what in skill_prompts(story, cp):
        doc.add_heading(skill, level=2)
        p = doc.add_paragraph()
        p.add_run("Ask: ").bold = True
        p.add_run(" · ".join(f"\"{q}\"" for q in prompts) or "(see the skill's description)")
        p = doc.add_paragraph()
        p.add_run("Expected response: ").bold = True
        p.add_run(what)
    doc.save(path)
    return path


# ── CLI ────────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="AiO Tableau Studio skills: generate story-loaded Studio Agent skills")
    ap.add_argument("command", choices=["story", "generate", "custom", "install", "status"])
    ap.add_argument("--slug", required=True)
    src = ap.add_mutually_exclusive_group()
    src.add_argument("--csv", help="story: bring an outside CSV into demos/{slug}/")
    src.add_argument("--hyper", help="story: bring an outside .hyper (or .tdsx with an extract)")
    src.add_argument("--datasource", help="story: download a published Tableau datasource by name (PAT)")
    ap.add_argument("--profile", help="config.json profile key for --datasource (default: active profile)")
    ap.add_argument("--new", action="store_true", help="story: ignore the saved story and re-derive it")
    ap.add_argument("--set", action="append", default=[], metavar="KEY=VALUE",
                    help="story: override a story key (value may be JSON)")
    ap.add_argument("--name", help="custom: kebab-case skill name")
    ap.add_argument("--spec", help="custom: path to the SKILL.md Claude wrote")
    ap.add_argument("--refs", help="custom: folder of extra reference files to include")
    a = ap.parse_args()

    if a.command == "story":
        if a.csv or a.hyper or a.datasource:
            ingest_source(a.slug, a.csv, a.hyper, a.datasource, a.profile)
        elif not os.path.exists(cp_path(a.slug)) and not glob.glob(os.path.join(demo_path(a.slug), "*.csv")):
            raise SystemExit("No checkpoint or CSV in this demo folder — pass --csv, --hyper, or --datasource.")
        story = build_story(a.slug, a.new, a.set)
        print(story_summary(story))
        print("\n" + json.dumps(story, indent=2, default=str))
    elif a.command == "generate":
        generate(a.slug)
    elif a.command == "custom":
        if not (a.name and a.spec):
            sys.exit("custom needs --name and --spec")
        add_custom_skill(a.slug, a.name, a.spec, a.refs)
    elif a.command == "install":
        install(a.slug)
    else:
        status(a.slug)


if __name__ == "__main__":
    main()
