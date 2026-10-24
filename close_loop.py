"""
Close-the-loop builder for AIO Analytics Builder.

Builds ONE follow-up action into the demo and the org, so the audience can trigger it during the meeting
(Claude Code is not running at demo time; the org closes the loop itself):

  AiO_{Slug}_Followup   autolaunched Flow — creates the story's records (Tasks per at-risk entity,
                         an escalation Case, or a recovery Campaign)
  AiO_{Slug}_Notify     record-triggered Flow — posts the "action taken" card to Slack (Chatter fallback)
  + ONE trigger, chosen from the story:
    slack_button        AiO_{Slug}_Alert (posts the insight card with a "Create follow-ups" link)
                        + AiO_{Slug}_Followup_Screen (the link target: confirm → run follow-up → done)
    record_triggered    AiO_{Slug}_Signal — a Task whose Subject starts with "AiO Signal:" runs the follow-up

CLI:
  python close_loop.py preflight   --slug <slug> [--profile key]
  python close_loop.py slack-check --slug <slug> [--profile key]   # read-only, JSON + fixes
  python close_loop.py plan        --slug <slug> [--action task|case|campaign] [--trigger slack_button|record_triggered] [--notify slack|chatter]
  python close_loop.py build       --slug <slug> [--action ...] [--trigger ...] [--notify slack|chatter] [--profile key] [--check-only]
  python close_loop.py test      --slug <slug> [--step trigger|followup]
  python close_loop.py reset     --slug <slug> [--yes]
  python close_loop.py remove    --slug <slug>
"""

import argparse
import ast
import base64
import glob
import io
import json
import os
import re
import sys
import time
import zipfile
from datetime import date, datetime, timedelta

import pandas as pd
import requests

ROOT = os.path.dirname(os.path.abspath(__file__))  # repo root, so demos/ resolves from any cwd

import close_loop_flows as flows

API = "v62.0"
SUBJECT_PREFIX = "AiO Alert"
SIGNAL_PREFIX = "AiO Signal"
SLACK_ACTION = "slackPostMessage"
TRIGGERS = ["slack_button", "record_triggered", "record_action"]

# Story keywords → action. Task is the default (path of least resistance).
CASE_KEYWORDS = ["stockout", "stock-out", "defect", "sla", "quality", "outage", "safety", "incident",
                 "downtime", "missed", "fill rate", "on-time", "on time", "compliance", "error",
                 "backlog", "delay", "route", "service level", "forecast accuracy", "mape"]
CAMPAIGN_KEYWORDS = ["participation", "enrollment", "enrolment", "adoption", "engagement", "attendance",
                     "redemption", "usage", "utilization", "opt-in", "uptake", "activation"]

# ── Checkpoint + story loading ─────────────────────────────────────────────────

def demo_dir(slug):
    d = os.path.join(ROOT, "demos", slug)
    if not os.path.isdir(d):
        raise SystemExit(f"Demo folder not found: {d}")
    return d


def cp_path(slug):
    return os.path.join(demo_dir(slug), f"{slug}_checkpoint.json")


def load_cp(slug):
    p = cp_path(slug)
    if os.path.exists(p):
        with open(p) as f:
            return json.load(f)
    return {}


def save_cp(slug, cp):
    with open(cp_path(slug), "w") as f:
        json.dump(cp, f, indent=2, default=str)


def _literal_assignments(script_path, names):
    """Read top-level literal assignments (COMPANY, METRIC_CONFIG, ...) without executing the script."""
    out = {}
    if not script_path or not os.path.exists(script_path):
        return out
    with open(script_path) as f:
        tree = ast.parse(f.read())
    for node in tree.body:
        if isinstance(node, ast.Assign):
            for t in node.targets:
                if isinstance(t, ast.Name) and t.id in names:
                    try:
                        out[t.id] = ast.literal_eval(node.value)
                    except ValueError:
                        pass  # non-literal (f-string, call) — skip
    return out


def _find_csv(slug, cp):
    p = cp.get("csv_path")
    if p and os.path.exists(p):
        return p
    story_ds = (cp.get("story") or {}).get("datasource")
    if story_ds:
        cand = story_ds if os.path.isabs(story_ds) else os.path.join(ROOT, story_ds)
        if os.path.exists(cand):
            return cand
    csvs = sorted(glob.glob(os.path.join(demo_dir(slug), "*.csv")), key=os.path.getsize, reverse=True)
    if not csvs:
        raise SystemExit(f"No CSV found in demos/{slug}/ — run the data phase first.")
    return csvs[0]  # largest CSV = the fact table


def load_story_context(slug):
    """Returns (cp, ctx) where ctx has company, use_case, metric_config, story, csv_path."""
    cp = load_cp(slug)
    scripts = glob.glob(os.path.join(demo_dir(slug), "*_demo.py"))
    lits = _literal_assignments(scripts[0] if scripts else None,
                                {"COMPANY", "USE_CASE_LABEL", "USE_CASE", "METRIC_CONFIG"})
    story = cp.get("story") or {}
    ctx = {
        "company": story.get("company") or lits.get("COMPANY") or slug.replace("_", " ").title(),
        "use_case": lits.get("USE_CASE_LABEL") or lits.get("USE_CASE") or "",
        "metric_config": lits.get("METRIC_CONFIG") or [],
        "story": story,
        "csv_path": _find_csv(slug, cp),
    }
    return cp, ctx


# ── Story analysis (computed from the actual data, never hand-written) ─────────

def _up_is_good(sentiment):
    s = str(sentiment or "").upper()
    return not ("DOWN_IS_GOOD" in s or s == "DOWN")


def _is_avg(mc):
    return "SUM" not in str(mc.get("pulse_agg", mc.get("agg", "AVERAGE"))).upper()


def _column_for(mc, columns):
    for key in ("field", "label"):
        if mc.get(key) in columns:
            return mc[key]
    return None


def _window_value(df, col, use_avg):
    """Mean of monthly values (monthly mean for rates, monthly sum for flows)."""
    if df.empty:
        return None
    monthly = df.groupby("_month")[col].mean() if use_avg else df.groupby("_month")[col].sum()
    return float(monthly.mean())


def _change(base_df, recent_df, col, use_avg):
    b, r = _window_value(base_df, col, use_avg), _window_value(recent_df, col, use_avg)
    if b is None or r is None or b == 0:
        return None
    return {"baseline": b, "recent": r, "pct": (r - b) / abs(b)}


def analyze_story(df, ctx, top_n=5):
    """Find the drill path the demo tells: overall → culprit → compound culprit → at-risk entities."""
    story = ctx["story"]
    cols = list(df.columns)
    date_field = story.get("date_field") or next((c for c in cols if c.lower() in ("date", "week", "month")), None)
    if not date_field:
        raise SystemExit("Could not find a Date column in the demo CSV.")
    df = df.copy()
    df[date_field] = pd.to_datetime(df[date_field])
    df["_month"] = df[date_field].dt.to_period("M")

    metrics = [m for m in ctx["metric_config"] if _column_for(m, cols)]
    if metrics:
        primary_mc = metrics[0]
    else:  # no METRIC_CONFIG — use the first non-target numeric column
        num = [c for c in df.select_dtypes("number").columns if "target" not in c.lower() and not c.startswith("_")]
        primary_mc = {"label": num[0], "field": num[0]}
    pcol = _column_for(primary_mc, cols)
    use_avg = _is_avg(primary_mc)
    up_good = _up_is_good(primary_mc.get("sentiment"))
    sign = -1 if up_good else 1  # "bad" score: positive = moving the wrong way

    # Entity = story's Q4 field, else highest-cardinality text column (excluding IDs)
    text_cols = [c for c in df.select_dtypes(exclude=["number", "datetime"]).columns
                 if c not in (date_field, "_month") and "record id" not in c.lower()]
    entity = story.get("entity_field") if story.get("entity_field") in cols else None
    if not entity:
        entity = max(text_cols, key=lambda c: df[c].nunique())
    dims = story.get("dimensions") or [c for c in text_cols if c != entity and 2 <= df[c].nunique() <= 30]
    dims = [d for d in dims if d in cols and d != entity]

    # Windows: recent = last 3 months; baseline = months at least 9 months before the latest
    latest = df["_month"].max()
    recent = df[df["_month"] > latest - 3]
    base = df[df["_month"] <= latest - 9]
    if base.empty:
        base = df[df["_month"] <= df["_month"].min() + (latest - df["_month"].min()).n // 2]

    def seg_changes(dim, b, r):
        out = []
        for val in sorted(r[dim].dropna().unique()):
            ch = _change(b[b[dim] == val], r[r[dim] == val], pcol, use_avg)
            if ch:
                out.append({"dimension": dim, "value": str(val), **ch, "bad": sign * ch["pct"]})
        return out

    overall = _change(base, recent, pcol, use_avg)
    by_dim = {d: seg_changes(d, base, recent) for d in dims}
    by_dim = {d: v for d, v in by_dim.items() if v}
    if not by_dim:
        raise SystemExit("No segment dimensions found to drill into.")

    # Culprit = the single segment moving the wrong way the most
    culprit = max((s for segs in by_dim.values() for s in segs), key=lambda s: s["bad"])
    counter = min(by_dim[culprit["dimension"]], key=lambda s: s["bad"])

    # Drill 2 — worst value of another dimension within the culprit
    mask_b = base[culprit["dimension"]] == culprit["value"]
    mask_r = recent[culprit["dimension"]] == culprit["value"]
    compound = None
    for d in dims:
        if d == culprit["dimension"]:
            continue
        for s in seg_changes(d, base[mask_b], recent[mask_r]):
            if compound is None or s["bad"] > compound["bad"]:
                compound = s
    root_dim = story.get("root_cause_dimension") or (compound or {}).get("dimension")

    # Supporting metrics within the culprit segment
    supporting = []
    wanted = story.get("supporting_metrics")  # optional: labels, in talk-track order
    support_mcs = ([m for lbl in wanted for m in metrics if m.get("label") == lbl] if wanted else metrics[1:4])
    for mc in support_mcs:
        col = _column_for(mc, cols)
        ch = _change(base[mask_b], recent[mask_r], col, _is_avg(mc))
        if ch:
            supporting.append({"label": mc["label"], **ch,
                               "wrong_way": (-1 if _up_is_good(mc.get("sentiment")) else 1) * ch["pct"] > 0})

    # At-risk = entities in the compound culprit (fallback: culprit) moving the wrong way the most
    scope_b, scope_r = base[mask_b], recent[mask_r]
    if compound:
        cb = scope_b[scope_b[compound["dimension"]] == compound["value"]]
        cr = scope_r[scope_r[compound["dimension"]] == compound["value"]]
        if cr[entity].nunique() >= top_n:
            scope_b, scope_r = cb, cr
    at_risk = []
    for ent in scope_r[entity].dropna().unique():
        ch = _change(scope_b[scope_b[entity] == ent], scope_r[scope_r[entity] == ent], pcol, use_avg)
        if ch and sign * ch["pct"] > 0:
            row = scope_r[scope_r[entity] == ent].iloc[-1]
            at_risk.append({"entity": str(ent), **ch,
                            "attributes": {d: str(row[d]) for d in dims if d in row}})
    at_risk.sort(key=lambda e: sign * (e["recent"] - e["baseline"]), reverse=True)  # absolute move

    return {
        "primary_metric": primary_mc.get("label", pcol),
        "primary_field": pcol,
        "is_rate": use_avg,
        "up_is_good": up_good,
        "entity_field": entity,
        "overall": overall,
        "culprit": culprit,
        "compound": compound,
        "counter_trend": counter if counter["value"] != culprit["value"] else None,
        "root_cause_dimension": root_dim,
        "supporting": supporting,
        "at_risk": at_risk[:top_n],
        "windows": {"recent_from": str(latest - 2), "recent_to": str(latest), "baseline_to": str(latest - 9)},
    }


# ── Action recommendation ──────────────────────────────────────────────────────

def recommend_action(ctx, analysis):
    """Pick the action that fits the story; Task is the simplest and the default."""
    text = " ".join([ctx.get("use_case", ""), analysis["primary_metric"],
                     ctx["story"].get("headline", "")]).lower()
    case_hits = [k for k in CASE_KEYWORDS if k in text]
    camp_hits = [k for k in CAMPAIGN_KEYWORDS if k in text]
    if camp_hits and len(camp_hits) >= len(case_hits):
        return "campaign", f"segment-wide behavior change ({', '.join(camp_hits)}) — fix is outreach to the whole segment"
    if case_hits:
        return "case", f"operational failure ({', '.join(case_hits)}) — one owner to fix the root cause"
    return "task", f"named at-risk {analysis['entity_field'].lower()}s — one follow-up per owner"


# ── Payload + message ──────────────────────────────────────────────────────────

def _fmt_val(v, is_rate):
    if is_rate and abs(v) <= 1.5:
        return f"{v * 100:.1f}%"
    return f"{v:,.2f}" if abs(v) < 100 else f"{v:,.0f}"


def _fmt_pct(p):
    return f"{'+' if p >= 0 else '−'}{abs(p) * 100:.0f}%"


def _fmt_ch(ch):
    """Relative % change, or point change when the metric crosses zero (e.g. NPS) where % is meaningless."""
    b, r = ch["baseline"], ch["recent"]
    if b * r < 0 or abs(ch["pct"]) > 3:
        d = r - b
        return f"{'+' if d >= 0 else '−'}{abs(d):,.1f} pts"
    return _fmt_pct(ch["pct"])


def _seg_label(s):
    return f"{s['value']} ({s['dimension']})"


def build_payload(ctx, analysis, action, reason, top_n=5, cfg=None, golden_record=None):
    a = analysis
    company, metric = ctx["company"], a["primary_metric"]
    culprit, compound = a["culprit"], a["compound"]
    focus = f"{culprit['value']} / {compound['value']}" if compound else culprit["value"]
    direction = "down" if a["overall"]["pct"] < 0 else "up"
    headline = f"{company}: {metric} {direction} {_fmt_ch(a['overall']).lstrip('+−')} — concentrated in {focus}"
    due = (date.today() + timedelta(days=2)).isoformat()

    drill = [f"Overall {_fmt_ch(a['overall'])}", f"{culprit['value']} {_fmt_ch(culprit)}"]
    if compound:
        drill.append(f"{culprit['value']} + {compound['value']} {_fmt_ch(compound)}")
    drill_path = " → ".join(drill)

    why = []
    if a["root_cause_dimension"]:
        why.append(f"Root cause shows up by {a['root_cause_dimension']}"
                   + (f" — {compound['value']} is the steepest" if compound and compound['dimension'] == a['root_cause_dimension'] else ""))
    for s in a["supporting"]:
        if s["wrong_way"]:
            why.append(f"{s['label']} {_fmt_ch(s)} in {culprit['value']}")
    counter = (f"{a['counter_trend']['value']} is {_fmt_ch(a['counter_trend'])} — this isn't systemic, it's {culprit['value']}."
               if a["counter_trend"] else "")

    at_risk_lines = [f"{e['entity']}: {metric} {_fmt_val(e['baseline'], a['is_rate'])} → "
                     f"{_fmt_val(e['recent'], a['is_rate'])} ({_fmt_ch(e)})" for e in a["at_risk"]]

    summary = "\n".join(filter(None, [
        headline, "", f"Where: {drill_path}",
        ("Why: " + "; ".join(why)) if why else "", counter, "",
        f"Most at risk ({a['entity_field']}):", *[f"- {l}" for l in at_risk_lines],
        "", f"Window: {a['windows']['recent_from']}–{a['windows']['recent_to']} vs. baseline through {a['windows']['baseline_to']}.",
        "Created by AiO Analytics Builder from the demo insight.",
    ]))

    # Golden record linking: if provided, link Tasks via WhatId, Cases via AccountId
    gr_id = golden_record.get("id") if golden_record else None
    gr_type = golden_record.get("type", "Account") if golden_record else None

    records = []
    if action == "task":
        for e, line in zip(a["at_risk"], at_risk_lines):
            rec = {"attributes": {"type": "Task"},
                   "Subject": f"{SUBJECT_PREFIX}: {metric} — {e['entity']}"[:255],
                   "Description": f"{line}\n\n{summary}"[:32000],
                   "ActivityDate": due, "Priority": "High", "Status": "Not Started"}
            if gr_id and gr_type in ("Account", "Opportunity", "Case"):
                rec["WhatId"] = gr_id
            records.append(rec)
    elif action == "case":
        rec = {"attributes": {"type": "Case"},
               "Subject": f"{SUBJECT_PREFIX}: {headline}"[:255],
               "Description": summary[:32000], "Priority": "High", "Status": "New"}
        if gr_id and gr_type == "Account":
            rec["AccountId"] = gr_id
        records.append(rec)
    elif action == "campaign":
        records.append({"attributes": {"type": "Campaign"},
                        "Name": f"AiO: {focus} {metric} recovery"[:80],
                        "Description": summary[:32000], "Status": "Planned", "IsActive": True,
                        "StartDate": date.today().isoformat(),
                        "EndDate": (date.today() + timedelta(days=30)).isoformat()})

    ask = {"task": f"Owners: review your follow-up Task by {due}.",
           "case": "Ops leads: the escalation Case is assigned — swarm on the root cause in the thread.",
           "campaign": f"Marketing/CS: the {focus} recovery Campaign is ready — add your accounts this week."}[action]

    return {
        "company": company, "metric": metric, "focus": focus, "action": action, "action_reason": reason,
        "headline": headline, "drill_path": drill_path, "why": why, "counter_trend": counter,
        "at_risk": a["at_risk"], "at_risk_lines": at_risk_lines, "summary": summary,
        "ask": ask, "records": records, "top_n": top_n,
        "slack": (cfg or {}).get("slack", {}), "analysis": a,
        "golden_record": golden_record,
        "staged_at": datetime.now().isoformat(timespec="seconds"),
    }


def _record_noun(action, n):
    return {"task": f"{n} follow-up Task{'s' if n != 1 else ''}", "case": "an escalation Case",
            "campaign": "a recovery Campaign"}[action]



def render_message(payload, mode="done", link_url=None, explore_url=None, slack=True):
    """The story card. mode="alert" ends with the Create-follow-ups link; mode="done" reports the action taken.
    Slack mrkdwn (slack=True) or plain text (Chatter)."""
    b = (lambda s: f"*{s}*") if slack else (lambda s: s)
    link = (lambda url, txt: f"<{url}|{txt}>") if slack else (lambda url, txt: f"{txt}: {url}")
    lines = [f"🔻 {b(payload['headline'])}", "",
             f"{b('Where:')} {payload['drill_path']}"]
    if payload["why"]:
        lines.append(f"{b('Why:')} " + "; ".join(payload["why"]))
    if payload["counter_trend"]:
        lines.append(f"{b('Not systemic:')} {payload['counter_trend']}")
    lines += ["", b(f"Most at risk ({payload['analysis']['entity_field']}):")]
    lines += [f"• {l}" for l in payload["at_risk_lines"]]
    noun = _record_noun(payload["action"], len(payload["records"]))
    if mode == "alert":
        if explore_url:
            lines += ["", f"{b('Explore:')} {link(explore_url, 'Open the metric')}"]
        lines += ["", f"👉 {link(link_url, 'Create follow-ups')} — creates {noun} in Salesforce."]
    else:
        lines += ["", f"{b('Action taken:')} Created {noun} in Salesforce."]
        if link_url:
            lines.append(f"• {link(link_url, 'Open in Salesforce')}")
        lines += ["", f"👉 {payload['ask']}"]
    return "\n".join(lines)


def explore_link(cp):
    """Best-effort link to the Pulse metric (Tableau Next dashboards have no stable public URL)."""
    try:
        from connections import load_config
        t = load_config().get("tableau", {})
        mids = cp.get("pulse_metric_ids") or []
        if isinstance(mids, dict):  # some demos key metric IDs by label
            mids = list(mids.values())
        if mids and t.get("server_url") and t.get("site_name"):
            return f"{t['server_url'].rstrip('/')}/pulse/site/{t['site_name']}/metrics/{mids[0]}"
    except Exception:
        pass
    return None


# ── Follow-up design: action × trigger → Flows ────────────────────────────────

def recommend_trigger(cp, action):
    """ONE trigger per demo. Golden record → record-page action; operational breaches → signal;
    everything else → Slack alert with a button."""
    if cp.get("golden_record"):
        return "record_action", (f"golden record selected ({cp['golden_record'].get('name', '?')}) "
                                 f"— a button on the record page is the most natural trigger")
    if action == "case":
        return "record_triggered", "operational threshold breach — the fix should fire automatically on the signal"
    why = "Pulse already alerts people — extend that alert into Slack with a one-click follow-up"
    if cp.get("next_done") or cp.get("sdm_done") or cp.get("next_sdm_done"):
        why = "Next build — Slack button until the Agentforce trigger is available"
    return "slack_button", why


def flow_names(slug, trigger):
    # Flow API names: letters/digits/single underscores, no trailing underscore, ≤ 80 chars
    base = "AiO_" + re.sub(r"[^A-Za-z0-9]+", "_", slug.title()).strip("_")[:28].rstrip("_")
    names = {"followup": f"{base}_Followup", "notify": f"{base}_Notify"}
    if trigger == "slack_button":
        names.update({"alert": f"{base}_Alert", "screen": f"{base}_Followup_Screen"})
    elif trigger == "record_triggered":
        names["signal"] = f"{base}_Signal"
    elif trigger == "record_action":
        names["action_screen"] = f"{base}_Action_Screen"
        names["quick_action"] = f"{base}_Action"
    return names


def signal_subject(payload):
    return f"{SIGNAL_PREFIX}: {payload['metric']} breach — {payload['focus']}"[:255]


def _flow_records(payload):
    """Payload records → Flow recordCreates (absolute dates become TODAY()+N formulas)."""
    offsets = {"ActivityDate": 2, "StartDate": 0, "EndDate": 30}
    out = []
    for rec in payload["records"]:
        fields = {k: v for k, v in rec.items() if k != "attributes" and k not in offsets}
        out.append({"sobject": rec["attributes"]["type"], "fields": fields,
                    "dates": {k: offsets[k] for k in rec if k in offsets}})
    return out


def _anchor(payload):
    """The record the Notify Flow listens for: the first record the follow-up creates (exact Subject/Name)."""
    rec = payload["records"][0]
    field = "Name" if rec["attributes"]["type"] == "Campaign" else "Subject"
    return rec["attributes"]["type"], field, rec[field]


def _post_spec(org, slack_cfg, notify=None):
    """notify: None = auto (Slack if ready, else Chatter) · "chatter" = always Chatter · "slack" = Slack or fail."""
    if notify == "chatter":
        return {"mode": "chatter"}
    if notify == "slack":
        diag = slack_diagnose(org, slack_cfg)
        if not diag["ready"]:
            lines = "\n".join(f"  - {i['problem']} → {i['fix']}" for i in diag["issues"])
            raise SystemExit(f"Slack was chosen but isn't ready:\n{lines}\n"
                             "Fix these (python close_loop.py slack-check --slug X) "
                             "or re-run with --notify chatter.")
    if org and org.get("slack_ready"):
        vals, missing = _slack_inputs(org["slack_inputs"], slack_cfg, "__MESSAGE__")
        if not missing:
            return {"mode": "slack", "action": org["slack_action"], "inputs": vals}
        print(f"⚠ Slack inputs missing {missing} — set them in config.json → slack. Using Chatter.")
    return {"mode": "chatter"}


def generate_flows(slug, payload, trigger, inst, post, explore_url=None):
    """Returns {flow_api_name: xml}."""
    n = flow_names(slug, trigger)
    company, metric = payload["company"], payload["metric"]
    obj, field, value = _anchor(payload)
    noun = _record_noun(payload["action"], len(payload["records"]))
    slack = post["mode"] == "slack"
    desc = f"AiO Analytics Builder close-the-loop for {company} — {payload['headline']}"
    out = {
        n["followup"]: flows.followup_flow(f"AiO {company} Follow-up", f"{desc}. Creates {noun}.",
                                           _flow_records(payload),
                                           golden_record=payload.get("golden_record")),
        n["notify"]: flows.notify_flow(
            f"AiO {company} Notify", f"{desc}. Posts the action-taken card when the follow-up runs.",
            obj, field, value,
            render_message(payload, "done", f"{inst}/lightning/r/{obj}/{{!$Record.Id}}/view", slack=slack), post),
    }
    if trigger == "slack_button":
        screen_url = f"{inst}/flow/{n['screen']}?retURL=/lightning/page/home"
        out[n["alert"]] = flows.alert_flow(
            f"AiO {company} Alert", f"{desc}. Posts the insight card with the Create-follow-ups link.",
            render_message(payload, "alert", screen_url, explore_url, slack=slack), post)
        confirm = (f"<p><b>{payload['headline']}</b></p><p>{payload['drill_path']}</p>"
                   f"<p>This creates {noun} in Salesforce and tells the team in "
                   f"{'Slack' if slack else 'Chatter'}.</p>")
        done = (f"<p><b>Done.</b> Created {noun}.</p>"
                f"<p><a href=\"/lightning/r/{obj}/{{!Run_Followup.firstRecordId}}/view\">Open the first record</a></p>")
        out[n["screen"]] = flows.screen_flow(f"AiO {company} Create Follow-ups", f"{desc}. Slack-button target.",
                                             n["followup"], confirm, done, "Create follow-ups")
    elif trigger == "record_triggered":
        out[n["signal"]] = flows.signal_flow(
            f"AiO {company} Signal", f"{desc}. A Task whose Subject starts with '{SIGNAL_PREFIX}: {metric}' "
            f"(a threshold breach) runs the follow-up automatically.",
            f"{SIGNAL_PREFIX}: {metric}"[:255], n["followup"])
    elif trigger == "record_action":
        gr = payload.get("golden_record") or {}
        gr_type = gr.get("type", "Account")
        gr_name = gr.get("name", "the record")
        confirm = (f"<p><b>{payload['headline']}</b></p><p>{payload['drill_path']}</p>"
                   f"<p>This creates {noun} linked to <b>{gr_name}</b> and tells the team in "
                   f"{'Slack' if slack else 'Chatter'}.</p>")
        done = (f"<p><b>Done.</b> Created {noun}.</p>"
                f"<p><a href=\"/lightning/r/{obj}/{{!Run_Followup.firstRecordId}}/view\">Open the first record</a></p>")
        out[n["action_screen"]] = flows.record_action_screen_flow(
            f"AiO {company} Action", f"{desc}. Record-page action — confirm → run follow-up → done.",
            n["followup"], confirm, done, "Create follow-ups")
        # QuickAction + RecordActionDeployment are non-Flow metadata — stored separately
        qa_member = f"{gr_type}.{n['quick_action']}"
        out[f"quickActions/{qa_member}.quickAction"] = flows.quick_action_xml(
            gr_type, n["quick_action"], n["action_screen"], f"AiO: {metric} Follow-up")
        deploy_name = f"AiO_{gr_type}_Actions"
        out[f"recordActionDeployments/{deploy_name}.recordActionDeployment"] = flows.record_action_deployment_xml(
            gr_type, qa_member, f"AiO {company} Actions")
        out["_quick_action_member"] = qa_member
        out["_deployment_name"] = deploy_name
    return out


def _talk_track(payload):
    """Story fragments the presenter speaks to: where it's concentrated, why, and the not-systemic proof."""
    a = payload["analysis"]
    worst = a["compound"] or a["culprit"]
    where = (f"{payload['focus'].replace(' / ', ' ')} is where it's worst — {payload['metric']} {_fmt_ch(worst)}, "
             f"against {_fmt_ch(a['overall'])} overall")
    why = next((w for w in payload["why"] if not w.startswith("Root cause")), None) or (payload["why"] or [None])[0]
    counter = (f"And it isn't company-wide: {a['counter_trend']['value']} is {_fmt_ch(a['counter_trend'])}."
               if a.get("counter_trend") else "")
    owner = {"task": f"each {a['entity_field'].lower()} owner", "case": "the ops lead who owns the Case",
             "campaign": "the team running the Campaign"}[payload["action"]]
    return where, why, counter, owner


def demo_script(slug, payload, trigger, names, slack_cfg, post_mode):
    """Section 7 steps: what to say, what to do (Trigger), what the audience sees (Result)."""
    chan = (slack_cfg or {}).get("channel_name") or (slack_cfg or {}).get("channel_id")
    where = (f"#{chan}" if chan else "the Slack channel") if post_mode == "slack" else "Chatter"
    noun = _record_noun(payload["action"], len(payload["records"]))
    focus, why, counter, owner = _talk_track(payload)
    because = f" The leading signal: {why}." if why else ""
    steps = []
    if trigger == "slack_button":
        steps.append({"title": "The alert lands",
                      "say": f"{focus}.{because} Most dashboards stop here. Let's see "
                             f"what happens when that insight lands where the team already works.",
                      "trigger": f"Before the meeting (or live), run the {names['alert']} Flow — `python close_loop.py "
                                 f"test --slug {slug}` — or open it in Setup → Flows → Run.",
                      "result": f"An insight card in {where}: {payload['headline']}, the drill path, and the at-risk "
                                f"list, with a \"Create follow-ups\" link."})
        steps.append({"title": "Close the loop",
                      "say": f"{counter} So what do we do? We assign it right now, from the alert, "
                             f"without opening another tool.".strip(),
                      "trigger": "Click \"Create follow-ups\" in the card and press the button on the confirmation "
                                 "screen.",
                      "result": f"Salesforce creates {noun} with the full story in each record, and a second card in "
                                f"{where} tells the team what was done and who owns it."})
    elif trigger == "record_action":
        gr = payload.get("golden_record") or {}
        gr_name = gr.get("name", "the record")
        gr_type = gr.get("type", "Account")
        steps.append({"title": "Navigate to the record",
                      "say": f"{focus}.{because} Now watch what happens when we go straight to the "
                             f"{gr_type} record — {gr_name}.",
                      "trigger": f"Open the {gr_type} record for {gr_name} in Salesforce.",
                      "result": f"The {gr_type} record page with the \"AiO: {payload['metric']} Follow-up\" "
                                f"action button visible in the Actions & Recommendations panel."})
        steps.append({"title": "Close the loop",
                      "say": f"{counter} We don't need to open another tool or remember to follow up. "
                             f"The action is right here on the record.".strip(),
                      "trigger": f"Click \"AiO: {payload['metric']} Follow-up\" in the Actions panel and press "
                                 f"\"Create follow-ups\" on the confirmation screen.",
                      "result": f"Salesforce creates {noun} linked to {gr_name}, and a notification in "
                                f"{where} tells the team what was done and who owns it."})
    else:
        steps.append({"title": "The signal fires",
                      "say": f"Nobody should have to spot this by hand. When {payload['metric']} breaches in "
                             f"{payload['focus']}, Salesforce opens the escalation on its own.",
                      "trigger": f"Log the threshold breach: create a Task with Subject \"{signal_subject(payload)}\" "
                                 f"(or let an upstream alert create it).",
                      "result": f"Salesforce automatically creates {noun} — nobody had to remember to escalate."})
        steps.append({"title": "The team is told",
                      "say": f"The team hears about it where they work, with the reason attached. {focus}.{because} "
                             f"{counter}".strip(),
                      "trigger": f"Open {where}.",
                      "result": f"An action-taken card with the drill path and a link to the record, ending: "
                                f"{payload['ask']}"})
    steps.append({"title": "Owners act",
                  "say": f"For {owner}, the evidence travels with the record, so nobody re-runs the analysis. "
                         f"That's insight to action in under a minute.",
                  "trigger": "Open one of the created records.",
                  "result": f"The record carries the evidence — {payload['drill_path']} — so the owner can act "
                            f"without re-running the analysis."})
    return steps


def section7(payload):
    """Walkthrough Section 7 content — shared by the .docx writer and the plan preview."""
    intro = (f"The demo ends with an action: {payload['action_reason']}. The follow-up is built into the org, so "
             f"nothing runs from a laptop at demo time. Trigger: {payload['trigger_label']}.")
    return intro, payload["demo_steps"]


# ── Golden record ─────────────────────────────────────────────────────────────

def lookup_golden_record(inst, h, record_type, search, limit=10):
    """Search for a CRM record by name or ID. Returns list of {Id, Name} dicts."""
    from crm_upserter import query_records
    search = (search or "").strip()
    if not search:
        return []
    if len(search) in (15, 18) and search[:3].isalnum():
        rows = query_records(inst.replace("Bearer ", ""), inst,
                             f"SELECT Id, Name FROM {record_type} WHERE Id = '{search}' LIMIT 1",
                             _headers=h)
        if rows:
            return rows
    safe = search.replace("'", "\\'")
    return query_records(None, None,
                         f"SELECT Id, Name FROM {record_type} WHERE Name LIKE '%{safe}%' "
                         f"ORDER BY LastModifiedDate DESC LIMIT {limit}",
                         _inst=inst, _headers=h)


def _lookup_soql(inst, h, soql):
    """Run a SOQL query and return records."""
    r = requests.get(f"{inst}/services/data/{API}/query",
                     headers=h, params={"q": soql}, timeout=30)
    if r.ok:
        return r.json().get("records", [])
    return []


def search_golden_record(inst, h, record_type, search_term, limit=10):
    """Search for a CRM record by name or ID. Returns list of {Id, Name} dicts."""
    search_term = (search_term or "").strip()
    if not search_term:
        return []
    if len(search_term) in (15, 18) and search_term[:3].isalnum():
        rows = _lookup_soql(inst, h,
            f"SELECT Id, Name FROM {record_type} WHERE Id = '{search_term}' LIMIT 1")
        if rows:
            return [{"Id": r["Id"], "Name": r.get("Name", "")} for r in rows]
    safe = search_term.replace("'", "\\'")
    rows = _lookup_soql(inst, h,
        f"SELECT Id, Name FROM {record_type} WHERE Name LIKE '%{safe}%' "
        f"ORDER BY LastModifiedDate DESC LIMIT {limit}")
    return [{"Id": r["Id"], "Name": r.get("Name", "")} for r in rows]


# ── Salesforce calls ───────────────────────────────────────────────────────────

def _sf(profile=None):
    from connections import get_sf_token, sf_headers
    token, inst = get_sf_token(profile_key=profile)
    return inst, sf_headers(token)


SLACK_FLOW_INPUTS = [  # slackPostMessage's required Flow inputs (found by check-only deploys)
    {"name": "slackAppIdForToken", "required": True, "label": "Slack App ID for Token"},
    {"name": "slackWorkspaceIdForToken", "required": True, "label": "Slack Workspace ID for Token"},
    {"name": "slackConversationId", "required": True, "label": "Slack Conversation ID"},
    {"name": "slackMessage", "required": True, "label": "Slack Message"},
]


def _slack_action_deployable(inst, h):
    """Check-only deploy of a one-step Flow calling slackPostMessage. Nothing is saved to the org."""
    params = "".join(f"<inputParameters><name>{p['name']}</name><value><stringValue>probe</stringValue>"
                     f"</value></inputParameters>" for p in SLACK_FLOW_INPUTS)
    flow = (f'<?xml version="1.0" encoding="UTF-8"?><Flow xmlns="http://soap.sforce.com/2006/04/metadata">'
            f"<actionCalls><name>Post</name><label>Post</label><locationX>0</locationX><locationY>0</locationY>"
            f"<actionName>{SLACK_ACTION}</actionName><actionType>{SLACK_ACTION}</actionType>{params}"
            f"<storeOutputAutomatically>true</storeOutputAutomatically></actionCalls>"
            f"<apiVersion>{API[1:]}</apiVersion><label>AiO Slack Probe</label><processType>AutoLaunchedFlow</processType>"
            f"<start><locationX>0</locationX><locationY>0</locationY><connector><targetReference>Post</targetReference>"
            f"</connector></start><status>Active</status></Flow>")
    try:
        deploy_metadata(inst, h, {"package.xml": flows.package_xml(["AiO_Slack_Probe"]),
                                  "flows/AiO_Slack_Probe.flow": flow}, check_only=True, timeout=120)
        return True
    except SystemExit:
        return False


def preflight(inst, h, slack_cfg=None):
    """Read-only org checks: object create perms, native Slack action + inputs, Metadata API reachable."""
    report = {"objects": {}, "slack_action": None, "slack_inputs": [], "slack_ready": False, "user_id": None}
    for obj in ("Task", "Case", "Campaign"):
        r = requests.get(f"{inst}/services/data/{API}/sobjects/{obj}/describe", headers=h, timeout=30)
        report["objects"][obj] = bool(r.ok and r.json().get("createable"))
    r = requests.get(f"{inst}/services/oauth2/userinfo", headers=h, timeout=30)
    if r.ok:
        report["user_id"] = r.json().get("user_id")
    r = requests.get(f"{inst}/services/data/{API}/actions/standard", headers=h, timeout=30)
    if r.ok:
        names = [a.get("name", "") for a in r.json().get("actions", [])]
        slack_names = [n for n in names if "slack" in n.lower()]
        report["slack_actions_available"] = slack_names
        name = SLACK_ACTION if SLACK_ACTION in names else next((n for n in slack_names if "message" in n.lower()), None)
        if name:
            report["slack_action"] = name
            d = requests.get(f"{inst}/services/data/{API}/actions/standard/{name}", headers=h, timeout=30)
            if d.ok:
                report["slack_inputs"] = [{"name": p["name"], "required": p.get("required", False),
                                           "label": p.get("label")} for p in d.json().get("inputs", [])]
    r = requests.get(f"{inst}/services/data/{API}/tooling/query", headers=h,
                     params={"q": "SELECT Id FROM FlowDefinition LIMIT 1"}, timeout=30)
    report["tooling_api"] = r.ok
    if not report["slack_action"] and _slack_action_deployable(inst, h):
        # Some orgs (e.g. Slack demo bundles) don't list slackPostMessage under /actions/standard,
        # but Flows can still use it — a check-only deploy of a probe Flow proves it.
        report["slack_action"] = SLACK_ACTION
        report["slack_inputs"] = SLACK_FLOW_INPUTS
    slack_cfg = slack_cfg or {}
    report["slack_ready"] = bool(report["slack_action"] and slack_cfg.get("channel_id"))
    return report


def _slack_inputs(action_inputs, slack_cfg, text):
    """Map our config onto the action's declared input names (names vary by release)."""
    vals = {}
    for p in action_inputs:
        n = p["name"].lower()
        if "message" in n or n.endswith("text"):
            vals[p["name"]] = text
        elif "conversation" in n or "channel" in n or "recipient" in n:
            vals[p["name"]] = slack_cfg.get("channel_id")
        elif "workspace" in n or "team" in n:
            vals[p["name"]] = slack_cfg.get("workspace_id")
        elif "appid" in n.replace("_", "") and slack_cfg.get("app_id"):
            vals[p["name"]] = slack_cfg["app_id"]
    missing = [p["name"] for p in action_inputs if p.get("required") and not vals.get(p["name"])]
    return vals, missing


def discover_slack_ids(inst, h, limit=40):
    """Read-only: find Slack app/workspace/channel IDs in the org's existing Active Flows that post to Slack
    (built in Flow Builder, so their IDs are known-good). Returns {} if none found."""
    try:
        rows = _tooling(inst, h, "SELECT Id FROM Flow WHERE Status = 'Active' "
                                 f"ORDER BY LastModifiedDate DESC LIMIT {int(limit)}")
    except requests.RequestException:
        return {}
    for row in rows:
        try:  # Metadata can only be selected one Flow at a time
            recs = _tooling(inst, h, f"SELECT Metadata FROM Flow WHERE Id = '{row['Id']}'")
        except requests.RequestException:
            continue
        for call in ((recs[0].get("Metadata") or {}).get("actionCalls") or []) if recs else []:
            if "slack" not in (call.get("actionType") or "").lower():
                continue
            found = {}
            for p in call.get("inputParameters") or []:
                n, v = p.get("name", "").lower(), ((p.get("value") or {}).get("stringValue") or "")
                if "appid" in n.replace("_", "") and v.startswith("A"):
                    found["app_id"] = v
                elif ("workspace" in n or "team" in n) and v.startswith("T"):
                    found["workspace_id"] = v
                elif ("conversation" in n or "channel" in n) and v[:1] in ("C", "G"):
                    found["channel_id"] = v
            if found:
                return found
    return {}


def slack_diagnose(org, slack_cfg, inst=None, h=None):
    """Turn a preflight report + saved Slack config into a ready flag and a list of plain-language fixes.
    With inst/h, also looks up known-good IDs from existing Slack Flows when something is missing."""
    slack_cfg = slack_cfg or {}
    issues = []

    def issue(code, problem, fix):
        issues.append({"code": code, "problem": problem, "fix": fix})

    if not (org or {}).get("slack_action"):
        issue("org_not_connected", "This Salesforce org has no Slack action — it isn't connected to Slack.",
              "In Salesforce: Setup → search 'Slack' → Slack Setup → accept the terms, then install the "
              "Salesforce for Slack (or Sales Cloud for Slack) app into the Slack workspace. Then, as the "
              "running user, connect your Slack account (open the Salesforce app in Slack → Home → Connect).")
    ch, ws, app = (slack_cfg.get(k, "") or "" for k in ("channel_id", "workspace_id", "app_id"))
    if not ch:
        issue("no_channel", "No Slack channel ID is saved.",
              "In Slack: open the channel → click its name → About tab → copy the Channel ID at the bottom "
              "(starts with C).")
    elif ch[0] not in ("C", "G"):
        issue("bad_channel", f"Channel ID '{ch}' doesn't look right (should start with C, or G for private).",
              "Copy the ID from the channel's About tab, not the channel name.")
    if not ws:
        issue("no_workspace", "No Slack workspace ID is saved.",
              "In Slack: workspace name menu → Settings → About this workspace, or copy the T… part of "
              "app.slack.com/client/T…/… in the browser.")
    elif ws[0] == "E":
        issue("enterprise_id", f"'{ws}' is an Enterprise Grid ID (E…), not a workspace ID.",
              "Use the workspace's own T… ID (shown in app.slack.com/client/T…/… when you open that workspace).")
    elif ws[0] != "T":
        issue("bad_workspace", f"Workspace ID '{ws}' should start with T (it's not the workspace name or URL).",
              "Copy the T… segment from app.slack.com/client/T…/….")
    if app and app[0] != "A":
        issue("bad_app_id", f"App ID '{app}' should start with A.", "Leave it blank unless the check asks for it.")
    if (org or {}).get("slack_action") and not issues:
        _, missing = _slack_inputs(org.get("slack_inputs", []), slack_cfg, "x")
        if missing:
            issue("missing_inputs", f"The org's Slack action also needs: {', '.join(missing)}.",
                  "Usually the Slack App ID (A…). The check looks it up from any existing Flow that posts to Slack; "
                  "otherwise build one Slack 'Send Message' Flow in Flow Builder once, then re-run the check.")
    if not (org or {}).get("tooling_api", True):
        issue("no_tooling", "The Tooling API isn't reachable for this user, so Flows can't be activated.",
              "Use a user with 'Modify All Data' / 'Manage Flow' (a System Administrator).")
    suggested = {}
    if inst and h and any(i["code"] in ("missing_inputs", "no_workspace", "bad_workspace", "enterprise_id",
                                        "no_channel") for i in issues) and (org or {}).get("slack_action"):
        suggested = discover_slack_ids(inst, h)
    return {"ready": not issues, "issues": issues, "suggested": suggested,
            "slack_action": (org or {}).get("slack_action"), "saved": {k: slack_cfg.get(k, "") for k in
                                                                     ("workspace_id", "channel_id", "channel_name", "app_id")}}


def _zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for path, xml in files.items():
            z.writestr(path, xml)
    return buf.getvalue()


def deploy_metadata(inst, h, files, check_only=False, timeout=300):
    """Zip {path: xml} and deploy via the REST Metadata API; poll until done. Raises on failure."""
    opts = {"deployOptions": {"checkOnly": check_only, "rollbackOnError": True, "singlePackage": True,
                              "ignoreWarnings": False}}
    r = requests.post(f"{inst}/services/data/{API}/metadata/deployRequest",
                      headers={"Authorization": h["Authorization"]},
                      files={"json": (None, json.dumps(opts), "application/json"),
                             "file": ("deploy.zip", _zip(files), "application/zip")},
                      timeout=30)
    if not r.ok:
        raise SystemExit(f"Deploy request failed ({r.status_code}): {r.text[:500]}")
    dep_id = r.json()["id"]
    t0 = time.time()
    while time.time() - t0 < timeout:
        time.sleep(3)
        s = requests.get(f"{inst}/services/data/{API}/metadata/deployRequest/{dep_id}",
                         headers=h, params={"includeDetails": "true"}, timeout=30)
        res = s.json().get("deployResult", {})
        if res.get("done"):
            if res.get("status") != "Succeeded":
                fails = (res.get("details") or {}).get("componentFailures") or []
                msgs = [f"{f.get('fullName')}: {f.get('problem')}" for f in fails] or [res.get("errorMessage")]
                raise SystemExit(f"Deploy {res.get('status')}:\n  " + "\n  ".join(map(str, msgs)))
            return dep_id
    raise SystemExit(f"Deploy {dep_id} still running after {timeout}s — check Setup → Deployment Status.")


def _tooling(inst, h, q):
    r = requests.get(f"{inst}/services/data/{API}/tooling/query", headers=h, params={"q": q}, timeout=30)
    r.raise_for_status()
    return r.json().get("records", [])


def set_flow_active(inst, h, names, active=True):
    """Activate the latest version (or deactivate) via Tooling FlowDefinition — production orgs deploy Flows
    inactive unless 'Deploy processes and flows as active' is on."""
    quoted = ",".join(f"'{n}'" for n in names)
    for d in _tooling(inst, h, f"SELECT Id, DeveloperName, ActiveVersionId, LatestVersionId "
                               f"FROM FlowDefinition WHERE DeveloperName IN ({quoted})"):
        if active and d["ActiveVersionId"] == d["LatestVersionId"]:
            continue
        if not active and not d["ActiveVersionId"]:
            continue
        num = 0
        if active:
            num = _tooling(inst, h, f"SELECT VersionNumber FROM Flow WHERE Id = '{d['LatestVersionId']}'")[0]["VersionNumber"]
        r = requests.patch(f"{inst}/services/data/{API}/tooling/sobjects/FlowDefinition/{d['Id']}", headers=h,
                           json={"Metadata": {"activeVersionNumber": num}}, timeout=30)
        print(f"{'✓' if r.ok else '⚠'} {'activated' if active else 'deactivated'} {d['DeveloperName']}"
              + ("" if r.ok else f" — {r.text[:200]}"))


def run_flow(inst, h, name):
    r = requests.post(f"{inst}/services/data/{API}/actions/custom/flow/{name}", headers=h, json={"inputs": [{}]}, timeout=30)
    body = r.json() if r.headers.get("content-type", "").startswith("application/json") else r.text
    ok = r.ok and isinstance(body, list) and body and body[0].get("isSuccess")
    if not ok:
        raise SystemExit(f"Flow {name} failed ({r.status_code}): {json.dumps(body)[:600]}")
    return body[0].get("outputValues") or {}


def _soql(inst, h, q):
    r = requests.get(f"{inst}/services/data/{API}/query", headers=h, params={"q": q}, timeout=30)
    r.raise_for_status()
    return r.json().get("records", [])


def _q(s):
    return "'" + str(s).replace("\\", "\\\\").replace("'", "\\'") + "'"


def find_followup_records(inst, h, payload, since=None):
    """Records the installed follow-up creates, matched on their exact Subject/Name (+ signal Tasks)."""
    found = []
    after = f" AND CreatedDate >= {since}" if since else ""
    by_obj = {}
    for rec in payload["records"]:
        obj = rec["attributes"]["type"]
        by_obj.setdefault(obj, []).append(rec.get("Subject") or rec.get("Name"))
    if payload.get("trigger") == "record_triggered":
        by_obj.setdefault("Task", []).append(signal_subject(payload))
    for obj, vals in by_obj.items():
        field = "Name" if obj == "Campaign" else "Subject"
        for r in _soql(inst, h, f"SELECT Id, {field} FROM {obj} WHERE {field} IN ({','.join(map(_q, vals))}){after}"):
            found.append({"sobject": obj, "id": r["Id"], "label": r[field]})
    return found


# ── Walkthrough Section 7 ──────────────────────────────────────────────────────

def write_walkthrough(slug, cp, payload):
    """Replace (or append) the 'Close the Loop' section in the demo's walkthrough (.docx or .html)."""
    # Find the walkthrough file — try .docx first, then .html
    path = cp.get("walkthrough")
    if not path or not os.path.exists(path):
        for ext in ("*.docx", "*.html"):
            candidates = sorted(glob.glob(os.path.join(demo_dir(slug), f"*walkthrough*{ext.lstrip('*')}")))
            candidates = [c for c in candidates if not os.path.basename(c).startswith("~$")]
            if candidates:
                path = candidates[0]
                break
    if not path or not os.path.exists(path):
        print("No walkthrough found — skipping Section 7 (talk track is in the payload JSON).")
        return None

    intro, steps = section7(payload)

    if path.endswith(".html"):
        return _write_walkthrough_html(path, intro, steps)
    else:
        return _write_walkthrough_docx(path, intro, steps)


def _write_walkthrough_docx(path, intro, steps):
    from docx import Document
    doc = Document(path)
    drop = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            drop = p.text.strip() == "Close the Loop"
        if drop:
            p._element.getparent().remove(p._element)

    doc.add_heading("Close the Loop", level=1)
    doc.add_paragraph(intro)
    for s in steps:
        doc.add_heading(s["title"], level=2)
        for label, key in (("Say: ", "say"), ("Trigger: ", "trigger"), ("Result: ", "result")):
            p = doc.add_paragraph()
            p.add_run(label).bold = True
            p.add_run(s[key])
    doc.save(path)
    return path


def _write_walkthrough_html(path, intro, steps):
    with open(path, "r") as f:
        html_content = f.read()

    # Remove existing Close the Loop section if present
    import re as _re
    html_content = _re.sub(
        r'<h1[^>]*>\s*Close the Loop\s*</h1>.*?(?=<h1[^>]*>|</body>|$)',
        '', html_content, flags=_re.DOTALL)

    # Build new section HTML
    section = ['<h1>Close the Loop</h1>', f'<p>{intro}</p>']
    for s in steps:
        section.append(f'<h2>{s["title"]}</h2>')
        section.append(f'<p><strong>Say:</strong> {s["say"]}</p>')
        section.append(f'<p><strong>Trigger:</strong> {s["trigger"]}</p>')
        section.append(f'<p><strong>Result:</strong> {s["result"]}</p>')
    section_html = '\n'.join(section)

    # Insert before </body> or append
    if '</body>' in html_content:
        html_content = html_content.replace('</body>', f'\n{section_html}\n</body>')
    else:
        html_content += f'\n{section_html}\n'

    with open(path, "w") as f:
        f.write(html_content)
    return path


# ── Commands ───────────────────────────────────────────────────────────────────

TRIGGER_LABELS = {"slack_button": "Slack message with a \"Create follow-ups\" button",
                  "record_triggered": "Record-triggered automation on a threshold-breach signal"}


def _cfg(profile=None):
    try:
        from connections import load_config
        return load_config(profile)
    except Exception:
        return {}


def _payload_path(slug):
    return os.path.join(demo_dir(slug), f"{slug}_close_loop_payload.json")


def _load_payload(slug):
    p = _payload_path(slug)
    if not os.path.exists(p):
        raise SystemExit(f"No close-loop plan — run `python close_loop.py build --slug {slug}` first.")
    with open(p) as f:
        return json.load(f)


def cmd_preflight(slug, profile=None):
    inst, h = _sf(profile)
    rep = preflight(inst, h, _cfg(profile).get("slack"))
    rep["instance"] = inst
    print(json.dumps(rep, indent=2))
    return rep


def cmd_slack_check(slug, profile=None):
    """Read-only: is Slack ready for this org + saved config? Prints JSON with plain-language fixes."""
    slack_cfg = _cfg(profile).get("slack", {})
    inst, h = _sf(profile)
    diag = slack_diagnose(preflight(inst, h, slack_cfg), slack_cfg, inst, h)
    diag["instance"] = inst
    print(json.dumps(diag, indent=2))
    return diag


def _saved_notify(slug):
    try:
        return load_cp(slug).get("close_loop", {}).get("notify")
    except Exception:
        return None


def cmd_plan(slug, action=None, trigger=None, top_n=5, org=None, profile=None, notify=None):
    """Analyze the story and design the ONE follow-up. Writes the payload JSON; no org changes."""
    cp, ctx = load_story_context(slug)
    cfg = _cfg(profile)
    analysis = analyze_story(pd.read_csv(ctx["csv_path"]), ctx, top_n)
    rec_action, reason = recommend_action(ctx, analysis)
    if action and action != rec_action:
        reason = f"overridden (recommended: {rec_action} — {reason})"
    action = action or rec_action
    if org and not org["objects"].get(action.capitalize()):
        print(f"⚠ {action.capitalize()} not createable for this user — falling back to Task.")
        action, reason = "task", f"fallback: {action.capitalize()} not createable"
    rec_trig, trig_reason = recommend_trigger(cp, action)
    if trigger and trigger != rec_trig:
        trig_reason = f"overridden (recommended: {rec_trig} — {trig_reason})"
    trigger = trigger or rec_trig

    payload = build_payload(ctx, analysis, action, reason, top_n, cfg)
    notify = notify or _saved_notify(slug)
    post_mode = "chatter" if notify == "chatter" else ("slack" if (org or {}).get("slack_ready") or notify == "slack"
                                                       else "chatter")
    names = flow_names(slug, trigger)
    payload.update({"trigger": trigger, "trigger_reason": trig_reason, "trigger_label": TRIGGER_LABELS[trigger],
                    "flows": names, "post_mode": post_mode if (org or notify) else "unknown (offline)"})
    payload["demo_steps"] = demo_script(slug, payload, trigger, names, cfg.get("slack", {}), post_mode)
    with open(_payload_path(slug), "w") as f:
        json.dump(payload, f, indent=2, default=str)

    print(f"Story:   {payload['headline']}")
    print(f"         {payload['drill_path']}")
    print(f"Action:  {action.upper()} ×{len(payload['records'])} — {reason}")
    print(f"Trigger: {TRIGGER_LABELS[trigger]} — {trig_reason}")
    print(f"Flows:   {', '.join(names.values())}")
    print(f"Notify:  {payload['post_mode']}")
    intro, steps = section7(payload)
    print("\n── Walkthrough Section 7 preview: Close the Loop (written to the .docx on build) ──")
    print(intro)
    for s in steps:
        print(f"\n• {s['title']}\n    Say:     {s['say']}\n    Trigger: {s['trigger']}\n    Result:  {s['result']}")
    print(f"\nPlan: {_payload_path(slug)}")
    return cp, payload


def cmd_build(slug, action=None, trigger=None, top_n=5, profile=None, check_only=False, notify=None):
    """Plan → generate Flows → deploy → activate → document. Tracks every deployed name in the checkpoint.
    notify: "slack" (fail with fixes if not ready) | "chatter" | None (reuse saved choice, else auto)."""
    cfg = _cfg(profile)
    notify = notify or _saved_notify(slug)
    inst, h = _sf(profile)
    org = preflight(inst, h, cfg.get("slack"))
    post = _post_spec(org, cfg.get("slack", {}), notify)  # before planning, so a Slack problem stops early
    cp, payload = cmd_plan(slug, action, trigger, top_n, org, profile, notify)
    xmls = generate_flows(slug, payload, payload["trigger"], inst, post, explore_link(cp))

    out_dir = os.path.join(demo_dir(slug), "close_loop_metadata")

    # Separate Flows from QuickActions/RecordActionDeployments
    flow_xmls = {}
    extra_files = {}
    qa_members = []
    deploy_names = []
    for name, xml in xmls.items():
        if name.startswith("quickActions/") or name.startswith("recordActionDeployments/"):
            extra_files[name] = xml
            if name.startswith("quickActions/"):
                qa_members.append(name.split("/")[1].replace(".quickAction", ""))
            elif name.startswith("recordActionDeployments/"):
                deploy_names.append(name.split("/")[1].replace(".recordActionDeployment", ""))
        elif name.startswith("_"):
            pass  # internal metadata keys, not files
        else:
            flow_xmls[name] = xml

    files = {"package.xml": flows.package_xml(list(flow_xmls), qa_members, deploy_names)}
    files.update({f"flows/{n}.flow": x for n, x in flow_xmls.items()})
    files.update(extra_files)
    for path, x in files.items():
        os.makedirs(os.path.dirname(os.path.join(out_dir, path)), exist_ok=True)
        with open(os.path.join(out_dir, path), "w") as f:
            f.write(x)
    print(f"\nMetadata written to {out_dir}")

    dep_id = deploy_metadata(inst, h, files, check_only=check_only)
    print(f"✓ Deploy {'validated' if check_only else 'succeeded'} ({dep_id})")
    if check_only:
        return
    cl = cp.get("close_loop", {})
    deployed = sorted(set(cl.get("deployed", [])) | set(xmls))
    cl.update({"deployed": deployed, "instance": inst, "profile": profile, "trigger": payload["trigger"],
               "notify": notify or cl.get("notify"),  # only an explicit choice is saved
               "action": payload["action"], "flows": payload["flows"], "payload_path": _payload_path(slug),
               "built_at": datetime.now().isoformat(timespec="seconds")})
    cp["close_loop"], cp["close_loop_built"] = cl, True
    cp.pop("close_loop_staged", None)
    save_cp(slug, cp)  # save before activation so remove can always find what was deployed

    set_flow_active(inst, h, list(xmls))
    doc = write_walkthrough(slug, cp, payload)
    if doc:
        print(f"✓ Walkthrough Section 7 updated: {doc}")
    print("\nInstalled. Rehearse with: python close_loop.py test --slug " + slug)


def cmd_test(slug, step="trigger"):
    """Fire the installed follow-up once, the same way the trigger would. Creates real records + posts."""
    cp = load_cp(slug)
    cl = cp.get("close_loop", {})
    if not cp.get("close_loop_built"):
        raise SystemExit("Nothing installed — run `close_loop.py build` first.")
    payload = _load_payload(slug)
    inst, h = _sf(cl.get("profile"))
    since = datetime.utcnow().replace(microsecond=0).isoformat() + "Z"
    names = cl["flows"]
    if step == "followup":
        print(f"Running {names['followup']} …")
        run_flow(inst, h, names["followup"])
    elif cl["trigger"] == "slack_button":
        print(f"Running {names['alert']} (posts the insight card) …")
        run_flow(inst, h, names["alert"])
        print("✓ Alert posted. Click \"Create follow-ups\" in the card to run the follow-up, "
              "then run `close_loop.py reset` when done rehearsing.")
    else:
        r = requests.post(f"{inst}/services/data/{API}/sobjects/Task", headers=h,
                          json={"Subject": signal_subject(payload), "Status": "Completed", "Priority": "High",
                                "Description": "AiO close-loop rehearsal signal."}, timeout=30)
        if not r.ok:
            raise SystemExit(f"Signal Task failed ({r.status_code}): {r.text[:400]}")
        print(f"✓ Signal Task created ({r.json()['id']}) — the Signal Flow runs the follow-up.")
    time.sleep(2)
    created = find_followup_records(inst, h, payload, since)
    tracked = {r["id"]: r for r in cl.get("test_records", [])}
    tracked.update({r["id"]: r for r in created})
    cl["test_records"] = list(tracked.values())
    cp["close_loop"] = cl
    save_cp(slug, cp)
    for r in created:
        print(f"✓ {r['sobject']}: {r['label']}\n  {inst}/lightning/r/{r['sobject']}/{r['id']}/view")
    if not created and step != "trigger":
        print("⚠ No follow-up records found yet — check Setup → Flows → Paused and Failed Flow Interviews.")


def cmd_reset(slug, yes=False):
    """Delete records the follow-up created. Checkpoint-tracked test records go without asking; records from
    live runs (same exact Subject/Name, created by the trigger in the meeting) are listed and need --yes."""
    cp = load_cp(slug)
    cl = cp.get("close_loop", {})
    payload = _load_payload(slug)
    inst, h = _sf(cl.get("profile"))
    tracked = {r["id"] for r in cl.get("test_records", [])}
    found = find_followup_records(inst, h, payload)
    untracked = [r for r in found if r["id"] not in tracked]
    ids = list(tracked)
    if untracked:
        print(f"{len(untracked)} matching record(s) not created by `test` (likely the live demo run):")
        for r in untracked:
            print(f"  {r['sobject']} {r['id']}: {r['label']}")
        if yes:
            ids += [r["id"] for r in untracked]
        else:
            print("Re-run with --yes to delete these too.")
    if not ids:
        print("Nothing to delete.")
    for i in range(0, len(ids), 200):
        r = requests.delete(f"{inst}/services/data/{API}/composite/sobjects",
                            headers=h, params={"ids": ",".join(ids[i:i + 200]), "allOrNone": "false"}, timeout=30)
        for res in (r.json() if r.ok else []):
            print(f"{'✓ deleted' if res.get('success') else '⚠ ' + str(res.get('errors'))} {res.get('id', '')}")
    print("Note: Slack/Chatter posts aren't deleted — remove them by hand in the channel.")
    cl["test_records"] = []
    cp["close_loop"] = cl
    save_cp(slug, cp)


def cmd_remove(slug):
    """Undeploy ONLY the Flows this tool deployed (cp.close_loop.deployed)."""
    cp = load_cp(slug)
    cl = cp.get("close_loop", {})
    names = cl.get("deployed", [])
    if not names:
        print("Nothing deployed for this demo.")
        return
    inst, h = _sf(cl.get("profile"))
    set_flow_active(inst, h, names, active=False)
    deploy_metadata(inst, h, {"package.xml": flows.package_xml(),
                              "destructiveChanges.xml": flows.package_xml(names)})
    print(f"✓ Removed {', '.join(names)}")
    cl["deployed"], cp["close_loop_built"] = [], False
    cp["close_loop"] = cl
    save_cp(slug, cp)


def main():
    ap = argparse.ArgumentParser(description="AiO close-the-loop: build ONE follow-up action into the demo org")
    ap.add_argument("command", choices=["preflight", "slack-check", "plan", "build", "test", "reset", "remove"])
    ap.add_argument("--slug", required=True)
    ap.add_argument("--action", choices=["task", "case", "campaign"])
    ap.add_argument("--trigger", choices=TRIGGERS)
    ap.add_argument("--top-n", type=int, default=5)
    ap.add_argument("--profile", help="config.json profile key (default: active profile)")
    ap.add_argument("--check-only", action="store_true", help="validate the deploy without saving it")
    ap.add_argument("--step", choices=["trigger", "followup"], default="trigger")
    ap.add_argument("--yes", action="store_true", help="reset: also delete untracked matching records")
    ap.add_argument("--notify", choices=["slack", "chatter"],
                    help="where the Notify Flow posts (default: saved choice, else Slack if ready, else Chatter)")
    a = ap.parse_args()
    if a.command == "preflight":
        cmd_preflight(a.slug, a.profile)
    elif a.command == "slack-check":
        cmd_slack_check(a.slug, a.profile)
    elif a.command == "plan":
        cmd_plan(a.slug, a.action, a.trigger, a.top_n, profile=a.profile, notify=a.notify)
    elif a.command == "build":
        cmd_build(a.slug, a.action, a.trigger, a.top_n, a.profile, a.check_only, a.notify)
    elif a.command == "test":
        cmd_test(a.slug, a.step)
    elif a.command == "reset":
        cmd_reset(a.slug, a.yes)
    else:
        cmd_remove(a.slug)


if __name__ == "__main__":
    main()
